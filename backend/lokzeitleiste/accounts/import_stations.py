"""Import an official DB station list into the local search catalog (CSV/XLSX)."""
import argparse
import csv
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from ..db import engine
from ..models import RailwayStation


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('file',type=Path)
    parser.add_argument('--code-column',default='Abk')
    parser.add_argument('--name-column',default='Name')
    parser.add_argument('--sheet')
    args=parser.parse_args()
    book=None
    if args.file.suffix.lower()=='.xlsx':
        book=load_workbook(args.file,read_only=True,data_only=True)
        sheet=book[args.sheet] if args.sheet else book.active
        rows=sheet.iter_rows(values_only=True)
    else:
        rows=csv.reader(args.file.read_text(encoding='utf-8-sig').splitlines(),delimiter=';')
    try:
        header=None
        for index,row in enumerate(rows):
            values=[str(c).strip() if c is not None else '' for c in row]
            if args.code_column in values and args.name_column in values:
                header=values;break
            if index>=50:break
        if header is None:raise SystemExit('Spalten nicht gefunden. --code-column und --name-column entsprechend der Datei setzen.')
        code_index=header.index(args.code_column);name_index=header.index(args.name_column)
        records={}
        for row in rows:
            if len(row)<=max(code_index,name_index):continue
            code=str(row[code_index] or '').strip().upper();name=' '.join(str(row[name_index] or '').split())
            if not code or not name:continue
            if len(code)>10 or len(name)>200:raise SystemExit('Ungültige Datensatzlänge; Import abgebrochen.')
            records[code]=name
        if not records:raise SystemExit('Keine Betriebsstellen gefunden.')
        with Session(engine()) as db:
            for code,name in records.items():
                row=db.get(RailwayStation,code)
                if row:row.name=name
                else:db.add(RailwayStation(code=code,name=name))
            db.commit()
        print(f'{len(records)} Betriebsstellen importiert/aktualisiert.')
    finally:
        if book:book.close()


if __name__=='__main__':main()
