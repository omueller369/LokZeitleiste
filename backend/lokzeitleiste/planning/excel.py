"""XLSX-Vertrag: Blatt Plan, Datum/Art/Notiz/Personalnummer; keine Formeln."""
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, GradientFill
from openpyxl.worksheet.datavalidation import DataValidation
from .schemas import DayInput
from ..i18n import translate, request_language

MAX_FILE = 2 * 1024 * 1024
MAX_UNPACKED = 20 * 1024 * 1024
COLORS = {"Arbeitstag": "DCEBF9", "Urlaub": "DBF1E3", "Ruhetag": "FFE8CC", "Ungeplant": "EEF0F3", "Feiertag":"E9DFF5"}


def parse_excel(content: bytes, *, year: int, personnel_number: str) -> list[DayInput]:
    if len(content) > MAX_FILE:
        raise ValueError("Excel-Datei darf höchstens 2 MB groß sein")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 500 or sum(x.file_size for x in archive.infolist()) > MAX_UNPACKED:
                raise ValueError("Excel-Datei ist entpackt zu groß")
            if any(x.filename.endswith("vbaProject.bin") for x in archive.infolist()):
                raise ValueError("Makros werden nicht importiert")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("Keine lesbare XLSX-Datei") from exc
    try:
        if "Plan" not in book.sheetnames:
            raise ValueError("Excel benötigt ein Blatt mit dem Namen Plan")
        sheet = book["Plan"]
        if (sheet.max_row or 0) > 367 or (sheet.max_column or 0) > 4:
            raise ValueError("Blatt Plan: höchstens 366 Tageszeilen und vier Spalten")
        sheet.reset_dimensions()
        rows = sheet.iter_rows(max_col=5, max_row=368)
        header_row = next(rows)
        if header_row[4].value not in (None, ""):
            raise ValueError("Blatt Plan benötigt höchstens vier Spalten")
        headers = [str(c.value or "").strip() for c in header_row[:4]]
        if headers[:2] != ["Datum", "Art"] or headers[2] not in ("", "Notiz") or headers[3] not in ("", "Personalnummer"):
            raise ValueError("Spalten: Datum, Art, optional Notiz, optional Personalnummer")
        result, seen = [], set()
        for number, row in enumerate(rows, start=2):
            values = [c.value for c in row]
            if all(v in (None, "") for v in values):
                continue
            if number > 367 or values[4] not in (None, ""):
                raise ValueError("Blatt Plan: höchstens 366 Tageszeilen und vier Spalten")
            values = values[:4]
            if any(c.data_type == "f" for c in row):
                raise ValueError(f"Zeile {number}: Formeln sind nicht erlaubt")
            raw_date, kind, note, person = values
            if isinstance(raw_date, datetime):
                day = raw_date.date()
            elif isinstance(raw_date, date):
                day = raw_date
            elif isinstance(raw_date, str):
                try:
                    day = date.fromisoformat(raw_date.strip()) if "-" in raw_date else datetime.strptime(raw_date.strip(), "%d.%m.%Y").date()
                except ValueError as exc:
                    raise ValueError(f"Zeile {number}: ungültiges Datum") from exc
            else:
                raise ValueError(f"Zeile {number}: Datum fehlt oder ist kein Excel-Datum")
            if day.year != year or day in seen:
                raise ValueError(f"Zeile {number}: falsches Jahr oder doppeltes Datum")
            if person is not None and str(person).strip() != personnel_number:
                raise ValueError(f"Zeile {number}: Personalnummer passt nicht zum ausgewählten Tf (als Text eingeben)")
            try:
                item = DayInput(date=day, kind=str(kind or "").strip(), note=str(note or ""))
            except ValueError as exc:
                raise ValueError(f"Zeile {number}: Art muss Arbeitstag, Urlaub, Ruhetag oder Ungeplant sein; Notiz höchstens 500 Zeichen") from exc
            result.append(item)
            seen.add(day)
        if not result:
            raise ValueError("Keine Tageszeilen gefunden")
        return result
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Excel-Inhalt ist beschädigt oder nicht unterstützt") from exc
    finally:
        book.close()


def template(year_plan: dict, personnel_number: str) -> bytes:
    book = Workbook()
    view = book.create_sheet(translate('Stundenansicht'))
    view.sheet_view.rightToLeft = request_language.get() == 'ar'
    view.append([translate(v) for v in ('Datum','Tagesart','Plansoll','Feiertag','Notiz')])
    for month in year_plan['months']:
        for item in month['days']:
            view.append([date.fromisoformat(item['date']),translate(item['kind']),item['target_minutes']/60,
                         translate(item['holiday_name']) if item.get('is_holiday') else '',item['note']])
            for cell in view[view.max_row]:
                if isinstance(cell.value,str):cell.data_type='s'
    for column in 'ABCDE':view.column_dimensions[column].width=25
    view.freeze_panes='A2'
    sheet = book.active
    sheet.title = "Plan"
    sheet.append(["Datum", "Art", "Notiz", "Personalnummer"])
    for month in year_plan["months"]:
        for item in month["days"]:
            sheet.append([date.fromisoformat(item["date"]), item["kind"], item["note"], personnel_number])
            row = sheet.max_row
            if item.get("is_holiday"):
                sheet.cell(row,1).comment=Comment("Feiertag Berlin: "+item["holiday_name"]+". Soll: 0 Stunden, unabhängig von der Tagesart.","LokZeitleiste")
            sheet.cell(row, 1).number_format = "DD.MM.YYYY"
            sheet.cell(row, 4).number_format = "@"
            # Nutzereingaben bleiben Text, auch wenn sie mit '=' beginnen.
            for column in (2, 3, 4):
                sheet.cell(row, column).data_type = "s"
            for cell in sheet[row]:
                if item.get('is_holiday') and item['kind']!='Ungeplant':
                    from openpyxl.styles.fills import Stop
                    cell.fill=GradientFill(type='linear',degree=45,stop=[Stop(COLORS['Feiertag'],0),Stop(COLORS['Feiertag'],0.499),Stop(COLORS[item['kind']],0.5),Stop(COLORS[item['kind']],1)])
                else:cell.fill = PatternFill("solid", fgColor=COLORS["Feiertag" if item.get("is_holiday") else item["kind"]])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="173E55")
        cell.font = Font(color="FFFFFF", bold=True)
    for col, width in {"A": 16, "B": 18, "C": 45, "D": 22}.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    validation = DataValidation(type="list", formula1='"Arbeitstag,Urlaub,Ruhetag,Ungeplant"', allow_blank=False)
    validation.errorTitle = "Ungültige Tagesart"
    validation.error = "Bitte einen Wert aus der Liste auswählen."
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(f"B2:B{sheet.max_row}")
    info=book.create_sheet('Feiertage Berlin')
    info.append(['Datum','Gesetzlicher Feiertag Berlin','Sollstunden'])
    for month in year_plan['months']:
        for item in month['days']:
            if item.get('is_holiday'):
                info.append([date.fromisoformat(item['date']),item['holiday_name'],0])
                info.cell(info.max_row,1).number_format='DD.MM.YYYY'
                info.cell(info.max_row,2).data_type='s'
                for cell in info[info.max_row]:cell.fill=PatternFill('solid',fgColor=COLORS['Feiertag'])
    info.column_dimensions['A'].width=16;info.column_dimensions['B'].width=65;info.column_dimensions['C'].width=16
    info.freeze_panes='A2'
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
