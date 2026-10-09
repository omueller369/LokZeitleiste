import io
import unittest
from datetime import date
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool
from lokzeitleiste.db import Base, database_session
from lokzeitleiste.main import app
from lokzeitleiste.models import User, TfProfile
from lokzeitleiste.security import hash_password


class PlanningTest(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite+pysqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(self.engine)
        def session():
            with Session(self.engine, expire_on_commit=False) as db:
                yield db
        app.dependency_overrides[database_session] = session
        with Session(self.engine) as db:
            admin = User(username="plan-admin", password_hash=hash_password("plan-admin-12345"), role="admin")
            tf = User(username="plan-tf", password_hash=hash_password("plan-tf-secret-12345"), role="tf")
            db.add_all([admin, tf]); db.flush(); self.tf_id = tf.id
            db.add(TfProfile(user_id=tf.id, last_name="Beispiel", first_name="Mira", personnel_number="00123",
                             target_hours_minutes=9600, vacation_days=30, birth_date=date(1990, 5, 12), bahncard=50,
                             email="mira@example.com", federal_state="BE")); db.commit()
        self.client = TestClient(app, base_url="https://testserver")
        self.client.post("/api/v1/admin/login", json={"username": "plan-admin", "password": "plan-admin-12345"})
        self.base = f"/api/v1/admin/tf/{self.tf_id}/plan"

    def tearDown(self):
        self.client.close(); app.dependency_overrides.clear(); self.engine.dispose()

    def put(self, days, month=9, year=2026, revision=0):
        return self.client.put(f"{self.base}/{year}/{month}", json={"expected_revision": revision, "days": days})

    def workbook(self, rows):
        book = Workbook(); sheet = book.active; sheet.title = "Plan"
        sheet.append(["Datum", "Art", "Notiz", "Personalnummer"])
        for row in rows: sheet.append(row)
        buffer = io.BytesIO(); book.save(buffer); return buffer.getvalue()

    def preview(self, rows, year=2026):
        return self.client.post(f"{self.base}/{year}/import/preview", files={"file": ("plan.xlsx", self.workbook(rows))})

    def test_days_totals_leapyear_and_revisions(self):
        self.assertEqual(self.client.get(f"{self.base}/2024/2").json()["unplanned_days"], 29)
        days = [{"date": "2026-09-01", "kind": "Arbeitstag"}, {"date": "2026-09-02", "kind": "Urlaub"}, {"date": "2026-09-03", "kind": "Ruhetag"}]
        saved = self.put(days).json()
        self.assertEqual((saved["work_days"], saved["vacation_days"], saved["rest_days"], saved["unplanned_days"]), (1, 1, 1, 27))
        self.assertEqual(saved["work_target_minutes"], 480); self.assertEqual(saved["target_minutes"], 960)
        self.assertEqual(self.put(days).status_code, 409)
        self.assertEqual(self.put([{"date": "2026-09-02", "kind": "Ungeplant"}], revision=1).status_code, 200)
        self.assertEqual(self.client.get(f"{self.base}/2026/9").json()["vacation_days"], 0)
        self.assertEqual(len(self.client.get(f"{self.base}/2026/history").json()), 4)
        self.assertEqual(self.put([{"date": "2026-10-01", "kind": "Urlaub"}], revision=2).status_code, 422)
        self.assertEqual(self.put([days[0], days[0]], revision=2).status_code, 422)

    def test_import_preview_atomic_commit_and_stale_revision(self):
        rows = [[date(2026, 9, 1), "Urlaub", "Sommerurlaub", "00123"], [date(2026, 10, 1), "Ruhetag", "", "00123"]]
        preview = self.preview(rows); self.assertEqual(preview.status_code, 200, preview.text)
        self.assertEqual(self.client.get(f"{self.base}/2026/9").json()["target_minutes"], 0)
        payload = {"days": [{k: d[k] for k in ("date", "kind", "note")} for d in preview.json()["days"]], "expected_revisions": preview.json()["expected_revisions"]}
        self.put([{"date": "2026-10-02", "kind": "Arbeitstag"}], month=10)
        self.assertEqual(self.client.post(f"{self.base}/2026/import", json=payload).status_code, 409)
        self.assertEqual(self.client.get(f"{self.base}/2026/9").json()["target_minutes"], 0)
        payload["expected_revisions"]["10"] = 1
        imported = self.client.post(f"{self.base}/2026/import", json=payload)
        self.assertEqual(imported.status_code, 200, imported.text); self.assertEqual(imported.json()["changed_days"], 2)
        self.assertEqual(self.client.get(f"{self.base}/2026/10").json()["work_days"], 1)

    def test_import_rejects_wrong_tf_formula_duplicates_and_invalid_file(self):
        invalid = [[[date(2026, 9, 1), "Urlaub", "", "someone-else"]], [[date(2026, 9, 1), "Urlaub", "=1+1", "00123"]],
                   [[date(2026, 9, 1), "Urlaub"], [date(2026, 9, 1), "Ruhetag"]], [[date(2027, 9, 1), "Urlaub"]], [[date(2026, 9, 1), "Unknown"]]]
        for rows in invalid:
            with self.subTest(rows=rows): self.assertEqual(self.preview(rows).status_code, 422)
        response = self.client.post(f"{self.base}/2026/import/preview", files={"file": ("bad.xlsx", b"bad")})
        self.assertEqual(response.status_code, 422); self.assertEqual(self.client.get(f"{self.base}/2026/9").json()["revision"], 0)

    def test_export_roundtrip_and_pdfs(self):
        self.put([{"date": "2024-02-29", "kind": "Urlaub", "note": "Erholung"}], month=2, year=2024)
        export = self.client.get(f"{self.base}/2024/template.xlsx"); self.assertEqual(export.status_code, 200)
        book = load_workbook(io.BytesIO(export.content)); sheet = book["Plan"]
        self.assertEqual(sheet.max_row, 367); self.assertEqual(sheet.cell(2, 4).value, "00123"); self.assertEqual(sheet.cell(2, 4).data_type, "s")
        response = self.client.post(f"{self.base}/2024/import/preview", files={"file": ("export.xlsx", export.content)})
        self.assertEqual(response.status_code, 200, response.text); self.assertEqual(len(response.json()["days"]), 366)
        for path in ["/2024/pdf", "/2024/2/pdf"]:
            pdf = self.client.get(self.base + path); self.assertEqual(pdf.status_code, 200); self.assertTrue(pdf.content.startswith(b"%PDF"))
        self.assertEqual(self.client.get(f"{self.base}/2024/13").status_code, 422)

    def test_plan_integrates_month_summary_without_double_vacation(self):
        days = [{"date": f"2026-09-{day:02d}", "kind": "Urlaub" if day == 1 else "Ruhetag"} for day in range(1, 31)]
        self.put(days)
        token = self.client.post("/api/v1/tf/login", json={"username": "plan-tf", "password": "plan-tf-secret-12345"}).json()["access_token"]
        summary = self.client.get("/api/v1/me/months/2026/9/summary", headers={"Authorization": "Bearer " + token}).json()
        self.assertEqual(summary["target_minutes"], 480); self.assertEqual(summary["balance_minutes"], -480); self.assertEqual(summary["totals"]["vacation"], 0)
        self.put([{"date": "2026-09-30", "kind": "Ungeplant"}], revision=1)
        summary = self.client.get("/api/v1/me/months/2026/9/summary", headers={"Authorization": "Bearer " + token}).json(); self.assertIsNone(summary["balance_minutes"])

    def test_admin_authorization_and_unknown_tf(self):
        other = TestClient(app, base_url="https://testserver")
        self.assertEqual(other.get(f"{self.base}/2026/9").status_code, 401)
        token = other.post("/api/v1/tf/login", json={"username": "plan-tf", "password": "plan-tf-secret-12345"}).json()["access_token"]
        self.assertEqual(other.get(f"{self.base}/2026/9", headers={"Authorization": "Bearer " + token}).status_code, 401)
        self.assertEqual(self.client.get("/api/v1/admin/tf/9999/plan/2026").status_code, 404); other.close()

    def test_berlin_calendar_2026_exact_and_exceptional_holidays(self):
        data=self.client.get(f'{self.base}/2026').json()
        days=[d for m in data['months'] for d in m['days'] if d['is_holiday']]
        expected={'2026-01-01','2026-03-08','2026-04-03','2026-04-06','2026-05-01','2026-05-14','2026-05-25','2026-10-03','2026-12-25','2026-12-26'}
        self.assertEqual({d['date'] for d in days},expected)
        self.assertEqual(data['totals']['holiday_days'],10)
        self.assertEqual(data['totals']['unplanned_days'],355)
        self.assertTrue(all(d['target_minutes']==0 and d['holiday_name'] and d['holiday_state']=='BE' for d in days))
        self.assertTrue(self.client.get(f'{self.base}/2025/5').json()['days'][7]['is_holiday'])
        self.assertFalse(self.client.get(f'{self.base}/2026/5').json()['days'][7]['is_holiday'])
        self.assertTrue(self.client.get(f'{self.base}/2028/6').json()['days'][16]['is_holiday'])
        self.assertFalse(self.client.get(f'{self.base}/2029/6').json()['days'][16]['is_holiday'])
        # Frauentag gilt erst seit 2019; Brandenburg-spezifischer Ostersonntag gilt nicht in Berlin.
        self.assertFalse(self.client.get(f'{self.base}/2018/3').json()['days'][7]['is_holiday'])
        self.assertFalse(self.client.get(f'{self.base}/2026/4').json()['days'][4]['is_holiday'])

    def test_holidays_override_work_vacation_and_do_not_leave_plan_open(self):
        days=[{'date':f'2026-05-{day:02d}','kind':'Arbeitstag'} for day in range(1,32)]
        data=self.put(days,month=5).json()
        self.assertTrue(data['complete']);self.assertEqual(data['holiday_days'],3)
        self.assertEqual(data['work_days'],28);self.assertEqual(data['target_minutes'],28*480)
        self.assertEqual(data['days'][0]['kind'],'Arbeitstag')
        self.assertEqual(data['days'][0]['target_minutes'],0)
        holiday=[{'date':'2026-05-01','kind':'Urlaub'},{'date':'2026-05-14','kind':'Ungeplant'},{'date':'2026-05-25','kind':'Ruhetag'}]
        data=self.put(holiday,month=5,revision=1).json()
        self.assertEqual(data['target_minutes'],28*480);self.assertEqual(data['vacation_days'],0)
        self.assertEqual(data['rest_days'],0);self.assertEqual(data['unplanned_days'],0);self.assertTrue(data['complete'])
        summary=self.client.get(f'/api/v1/admin/tf/{self.tf_id}/months/2026/5/summary').json()
        self.assertEqual(summary['target_minutes'],28*480)
        self.assertEqual(summary['balance_minutes'],-28*480)
        self.assertEqual(summary['totals']['credited'],0)

    def test_holidays_excel_preview_bulk_and_export_keep_zero_target(self):
        preview=self.preview([[date(2026,5,1),'Arbeitstag','Bestehender Dienst','00123'],[date(2026,5,2),'Urlaub','','00123']]).json()
        self.assertTrue(preview['days'][0]['is_holiday']);self.assertEqual(preview['days'][0]['target_minutes'],0)
        payload={'days':[{k:d[k] for k in ('date','kind','note')} for d in preview['days']], 'expected_revisions':preview['expected_revisions']}
        data=self.client.post(f'{self.base}/2026/import',json=payload).json()['plan']['months'][4]
        self.assertEqual(data['target_minutes'],480);self.assertEqual(data['work_days'],0)
        self.assertEqual(self.client.post(f'{self.base}/2026/bulk',json={'days':[{'date':'2026-05-01','kind':'Urlaub'}],'expected_revisions':{'5':1}}).status_code,200)
        data=self.client.get(f'{self.base}/2026/5').json();self.assertEqual(data['target_minutes'],480);self.assertEqual(data['vacation_days'],1)
        content=self.client.get(f'{self.base}/2026/template.xlsx').content
        book=load_workbook(io.BytesIO(content));sheet=book['Plan']
        row=next(r for r in sheet.iter_rows(min_row=2) if r[0].value.date()==date(2026,5,1))
        self.assertEqual(row[0].fill.fgColor.rgb,'00E9DFF5');self.assertIn('Soll: 0 Stunden',row[0].comment.text)
        self.assertEqual(row[1].value,'Urlaub');self.assertEqual(row[2].value,'') if row[2].value=='' else self.assertIsNone(row[2].value)
        self.assertEqual(self.client.post(f'{self.base}/2026/import/preview',files={'file':('plan.xlsx',content)}).status_code,200)
